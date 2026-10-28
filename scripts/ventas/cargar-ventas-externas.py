#!/usr/bin/env python
"""
CARGA LAS VENTAS QUE NO PASARON POR SAGE
========================================

De dónde salen: el cuadro de dirección «PRESUPUESTO ... CIFRAS VENTAS AGENTE
VENDEDOR.xlsx», hoja «6,15», que lleva mes a mes el real de cada empresa del
grupo, incluidas las que entonces se llevaban fuera de Sage.

Qué carga, y por qué solo eso:
  · Jender, de enero a septiembre de 2025. En Sage entra el 6/10/2025, así que
    de octubre en adelante ya está y volver a sumarlo sería contarlo dos veces.
    Comprobado: octubre a diciembre da 40.838 € en el Excel y 39.782 € en
    Sage — es el mismo dato con sus abonos, no dos cosas distintas.

Y nada más. En particular NO carga Blizcool, aunque el Excel tenga su columna:
lo que ahí se llama Blizcool es en Sage la sociedad Tools Place, y está dentro
desde febrero de 2026. Se cargó una vez y sumó 253.264 € de más al año 2026
—los cuatro primeros meses coinciden al céntimo con Sage—, así que si algún día
hace falta volver a meterlo, primero hay que comprobar qué sociedad lo lleva.

No inventa nada: un mes sin cifra en el Excel no se carga.

Solo escribe en external_sales. No toca ninguna tabla de Sage, y volver a
ejecutarlo no duplica: cada origen y mes es una fila única.

Uso (desde la raíz del proyecto, con .env.local puesto):
    python scripts/ventas/cargar-ventas-externas.py --excel "C:\\ruta\\al.xlsx"
    python scripts/ventas/cargar-ventas-externas.py --excel "..." --seco
"""
import argparse
import json
import sys
import urllib.error
import urllib.request
from pathlib import Path

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
         "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

"""
Qué columna del Excel es cada cosa.

Van por número de columna y no buscando el texto de la cabecera a propósito:
ese cuadro reparte los títulos en cinco filas con celdas combinadas, y
buscarlos por texto acierta hoy y falla el día que alguien mueva una. El
número se comprueba contra lo que se espera leer encima, y si no cuadra el
cargador para en vez de meter la columna de al lado.
"""
BLOQUES = [
    {
        "fuente": "Jender",
        "sociedad": 2,
        # La ficha de comercial que firma esa venta en Sage de octubre en
        # adelante, y la que lleva el bloque en el cuadro. Va por código y no
        # por nombre: este repositorio es público.
        "comercial": 136,
        "columna": 19,                       # «REAL 2025» del bloque JENDER
        "espera_encima": ["JENDER", "REAL 2025"],
        "anio": 2025,
        "meses": [1, 2, 3, 4, 5, 6, 7, 8, 9],
        "nota": "Cuadro de direccion, hoja 6,15. Jender entro en Sage el 6/10/2025.",
    },
]


def variables_de_entorno():
    """Lo que hace falta de .env.local. La clave no se imprime nunca."""
    env = {}
    for linea in Path(".env.local").read_text(encoding="utf-8").splitlines():
        if "=" not in linea or linea.strip().startswith("#"):
            continue
        clave, valor = linea.split("=", 1)
        env[clave.strip()] = valor.strip()
    if not env.get("NEXT_PUBLIC_SUPABASE_URL") or not env.get("SUPABASE_SERVICE_ROLE_KEY"):
        print("Faltan NEXT_PUBLIC_SUPABASE_URL o SUPABASE_SERVICE_ROLE_KEY en .env.local.")
        sys.exit(1)
    return env


def euros(valor):
    return f"{valor:,.2f} EUR".replace(",", "X").replace(".", ",").replace("X", ".")


def leer(ruta):
    """Los meses de cada bloque, tal como están en el Excel."""
    import openpyxl

    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    if "6,15" not in libro.sheetnames:
        print(f"El Excel no tiene la hoja «6,15». Tiene: {', '.join(libro.sheetnames)}")
        sys.exit(1)
    filas = list(libro["6,15"].iter_rows(values_only=True))
    libro.close()

    def celda(fila, col):
        return fila[col] if col < len(fila) else None

    salida = []
    for bloque in BLOQUES:
        col = bloque["columna"] - 1
        # El nombre del bloque vive en una celda combinada a la izquierda de
        # su columna —«JENDER» encabeza presupuesto, real y desviación—, así
        # que se mira esa y las dos de al lado, no solo la suya.
        encima = " ".join(
            str(celda(filas[i], c)).strip()
            for i in range(0, min(5, len(filas)))
            for c in range(max(col - 2, 0), col + 2)
            if celda(filas[i], c) is not None
        ).lower()
        if not all(parte.lower() in encima for parte in bloque["espera_encima"]):
            print(f"La columna {bloque['columna']} no es la de {bloque['fuente']}: encima pone «{encima[:60]}».")
            print("El Excel ha cambiado de forma. Revisa las columnas antes de cargar nada.")
            sys.exit(1)

        for fila in filas:
            etiqueta = str(celda(fila, 1) or "").strip().upper()
            if etiqueta not in MESES:
                continue
            mes = MESES.index(etiqueta) + 1
            if mes not in bloque["meses"]:
                continue
            valor = celda(fila, col)
            if not isinstance(valor, (int, float)) or not valor:
                continue
            salida.append({
                "source": bloque["fuente"],
                "company_code": bloque["sociedad"],
                "rep_code": bloque["comercial"],
                "month": f"{bloque['anio']}-{mes:02d}-01",
                "net_amount": round(float(valor), 2),
                "note": bloque["nota"],
            })
    return salida


def cargar(env, filas):
    cuerpo = json.dumps(filas).encode("utf-8")
    peticion = urllib.request.Request(
        f"{env['NEXT_PUBLIC_SUPABASE_URL'].rstrip('/')}/rest/v1/external_sales?on_conflict=source,month",
        data=cuerpo,
        method="POST",
        headers={
            "apikey": env["SUPABASE_SERVICE_ROLE_KEY"],
            "Authorization": f"Bearer {env['SUPABASE_SERVICE_ROLE_KEY']}",
            "Content-Type": "application/json",
            "Prefer": "resolution=merge-duplicates,return=minimal",
        },
    )
    try:
        with urllib.request.urlopen(peticion, timeout=60) as respuesta:
            return respuesta.status
    except urllib.error.HTTPError as fallo:
        # El cuerpo del error dice qué columna no cuadra; la clave no sale ahí.
        print(f"No se pudo cargar ({fallo.code}): {fallo.read().decode('utf-8', 'replace')[:300]}")
        sys.exit(1)


def main():
    p = argparse.ArgumentParser(description="Carga en el Hub las ventas que no pasaron por Sage.")
    p.add_argument("--excel", required=True, help="el cuadro de dirección, en .xlsx")
    p.add_argument("--seco", action="store_true", help="leer y enseñar, sin escribir nada")
    args = p.parse_args()

    if not Path(args.excel).is_file():
        print(f"No encuentro el Excel:\n  {args.excel}")
        return 1

    filas = leer(args.excel)
    if not filas:
        print("No hay ni un mes con cifra. Revisa el Excel.")
        return 1

    print(f"Se van a cargar {len(filas)} meses:")
    for fuente in dict.fromkeys(fila["source"] for fila in filas):
        suyas = [fila for fila in filas if fila["source"] == fuente]
        total = sum(fila["net_amount"] for fila in suyas)
        print(f"  {fuente}: {len(suyas)} meses, {euros(total)}"
              f" ({suyas[0]['month'][:7]} -> {suyas[-1]['month'][:7]})")

    if args.seco:
        print("\n(en seco: no se ha escrito nada)")
        return 0

    cargar(variables_de_entorno(), filas)
    print("\nCargado. Volver a ejecutarlo no duplica nada: cada origen y mes es una fila.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
